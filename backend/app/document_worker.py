"""Bounded document parser worker; invoked in a fresh interpreter per document.

This is process/resource isolation, not a complete operating-system security
sandbox. No files are extracted from archives and no application configuration
or credentials are imported.
"""

import json
import os
import socket
import sys
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZipFile


MAX_INPUT_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1000
MAX_PDF_PAGES = 200
MAX_OUTPUT_CHARACTERS = 250_000
MEMORY_LIMIT_BYTES = 384 * 1024 * 1024
CPU_LIMIT_SECONDS = 6


class DocumentLimitError(ValueError):
    """A document exceeds an explicitly supported resource budget."""


class _TextParts:
    def __init__(self, separator: str):
        self.parts = []
        self.characters = 0
        self.separator = separator

    def append(self, text: str):
        total = self.characters + len(text)
        if self.parts:
            total += len(self.separator)
        if total > MAX_OUTPUT_CHARACTERS:
            raise DocumentLimitError("Document text limit exceeded")
        self.parts.append(text)
        self.characters = total

    def text(self):
        return self.separator.join(self.parts)


def _check_docx_archive(file_bytes: bytes):
    with ZipFile(BytesIO(file_bytes)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise DocumentLimitError("Document archive entry limit exceeded")
        if sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
            raise DocumentLimitError("Document expanded content limit exceeded")
        expanded = 0
        for entry in entries:
            path = PurePosixPath(entry.filename)
            if path.is_absolute() or ".." in path.parts or entry.flag_bits & 1:
                raise ValueError("Invalid document archive")
            # Verify actual decoded bytes too; never extract archive paths.
            with archive.open(entry) as stream:
                tail = b""
                while chunk := stream.read(64 * 1024):
                    expanded += len(chunk)
                    if expanded > MAX_EXPANDED_BYTES:
                        raise DocumentLimitError("Document expanded content limit exceeded")
                    if entry.filename.lower().endswith((".xml", ".rels")):
                        scan = (tail + chunk).upper()
                        if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
                            raise ValueError("Document XML declarations are not supported")
                        tail = scan[-16:]


def extract_docx_text(file_bytes: bytes) -> str:
    from docx import Document

    _check_docx_archive(file_bytes)
    document = Document(BytesIO(file_bytes))
    parts = _TextParts("\n")
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    row_text.append(text)
            if row_text:
                parts.append(" | ".join(row_text))
    return parts.text()


def extract_pdf_text(file_bytes: bytes) -> str:
    from pypdf import apply_configuration
    from pypdf.errors import LimitReachedError

    try:
        # Apply library bounds before decoding streams, including nested form
        # resources. Total page content is separately counted below; the worker
        # address-space/CPU limits cover other parser allocations and work.
        with apply_configuration(
            maximum_declared_stream_length=MAX_EXPANDED_BYTES,
            array_based_stream_maximum_output_length=MAX_EXPANDED_BYTES,
            jbig2_maximum_output_length=MAX_EXPANDED_BYTES,
            lzw_maximum_output_length=MAX_EXPANDED_BYTES,
            run_length_maximum_output_length=MAX_EXPANDED_BYTES,
            zlib_maximum_output_length=MAX_EXPANDED_BYTES,
            brotli_maximum_output_length=MAX_EXPANDED_BYTES,
            image_maximum_buffer_size=MAX_EXPANDED_BYTES,
            stream_decoding_work_maximum_length=MAX_EXPANDED_BYTES,
            jbig2dec_binary=None,
        ):
            return _extract_pdf_text(file_bytes)
    except LimitReachedError:
        raise DocumentLimitError("Document expanded content limit exceeded") from None


def _extract_pdf_text(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(file_bytes))
    if reader.is_encrypted:
        raise ValueError("Encrypted documents are not supported")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentLimitError("Document page limit exceeded")
    expanded = 0
    parts = _TextParts("\n\n")
    for page in reader.pages:
        contents = page.get_contents()
        if contents is not None:
            expanded += len(contents.get_data())
            if expanded > MAX_EXPANDED_BYTES:
                raise DocumentLimitError("Document expanded content limit exceeded")
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())
    return parts.text()


def _disable_network():
    def blocked(*args, **kwargs):
        raise OSError("Document parser network access is disabled")

    socket.socket = blocked
    socket.SocketType = blocked
    socket.create_connection = blocked
    socket.getaddrinfo = blocked


def _apply_resource_limits():
    # Linux/Render supports these limits. Fail closed on unsupported platforms
    # rather than silently parsing an untrusted document without the budget.
    import resource

    resource.setrlimit(resource.RLIMIT_AS, (MEMORY_LIMIT_BYTES, MEMORY_LIMIT_BYTES))
    resource.setrlimit(resource.RLIMIT_CPU, (CPU_LIMIT_SECONDS, CPU_LIMIT_SECONDS + 1))
    resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
    resource.setrlimit(resource.RLIMIT_NOFILE, (32, 32))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    os.umask(0o077)


def main() -> int:
    wire_output = sys.stdout
    # Parser diagnostics are not part of the bounded IPC protocol.
    sys.stdout = sys.stderr
    try:
        _apply_resource_limits()
        _disable_network()
        data = sys.stdin.buffer.read(MAX_INPUT_BYTES + 1)
        if len(data) > MAX_INPUT_BYTES:
            raise DocumentLimitError("Document size limit exceeded")
        extension = sys.argv[1]
        if extension == ".pdf":
            text = extract_pdf_text(data)
        elif extension == ".docx":
            text = extract_docx_text(data)
        else:
            raise ValueError("Unsupported document")
        response = {"ok": True, "text": text}
    except (DocumentLimitError, MemoryError):
        response = {"ok": False, "kind": "limit"}
    except (ImportError, OSError):
        # Distinguish inability to establish the sandbox/import the parser from
        # document validation failures without reflecting exception details.
        response = {"ok": False, "kind": "unavailable"}
    except Exception:
        response = {"ok": False, "kind": "invalid"}
    wire_output.write(json.dumps(response, ensure_ascii=True, separators=(",", ":")))
    wire_output.flush()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
