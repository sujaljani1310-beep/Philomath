import asyncio
import os
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import HTTPException, UploadFile

from app import document_extraction as extraction
from app import document_worker as worker
from app.routes import files
from app.services.auth_service import AuthenticatedUser


def docx_bytes(text="Hello document"):
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_bytes(pages=1, text="Hello PDF"):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class ExtractionContentTests(unittest.TestCase):
    def test_normal_documents_preserve_text_and_tables(self):
        self.assertEqual(worker.extract_docx_text(docx_bytes("  Hello document  ")), "Hello document\nleft | right")
        self.assertEqual(worker.extract_pdf_text(pdf_bytes(2)), "Hello PDF\n\nHello PDF")

    def test_docx_expansion_and_entry_limits_and_invalid_xml(self):
        for name, value, expected in (
            ("expanded", b"x" * 2048, worker.DocumentLimitError),
            ("declaration.xml", b'<!DOCTYPE x [<!ENTITY y "boom">]><x/>', ValueError),
        ):
            with self.subTest(name=name):
                stream = BytesIO()
                with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
                    archive.writestr(name, value)
                with patch.object(worker, "MAX_EXPANDED_BYTES", 1024):
                    with self.assertRaises(expected):
                        worker._check_docx_archive(stream.getvalue())
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            archive.writestr("first", b"a")
            archive.writestr("second", b"b")
        with patch.object(worker, "MAX_ARCHIVE_ENTRIES", 1):
            with self.assertRaises(worker.DocumentLimitError):
                worker._check_docx_archive(stream.getvalue())

    def test_page_expansion_and_output_limits(self):
        with patch.object(worker, "MAX_PDF_PAGES", 1):
            with self.assertRaises(worker.DocumentLimitError):
                worker.extract_pdf_text(pdf_bytes(2))
        with patch.object(worker, "MAX_EXPANDED_BYTES", 1):
            with self.assertRaises(worker.DocumentLimitError):
                worker.extract_pdf_text(pdf_bytes())
        with patch.object(worker, "MAX_OUTPUT_CHARACTERS", 8):
            for extract, content in ((worker.extract_pdf_text, pdf_bytes()), (worker.extract_docx_text, docx_bytes())):
                with self.assertRaises(worker.DocumentLimitError):
                    extract(content)
        with patch.object(worker, "MAX_OUTPUT_CHARACTERS", len("Hello PDF")):
            self.assertEqual(worker.extract_pdf_text(pdf_bytes()), "Hello PDF")

    def test_os_resource_limits_are_applied(self):
        import resource

        with patch.object(resource, "setrlimit") as set_limit, patch.object(os, "umask"):
            worker._apply_resource_limits()
        limits = dict(call.args for call in set_limit.call_args_list)
        self.assertEqual(limits[resource.RLIMIT_AS], (384 * 1024 * 1024,) * 2)
        self.assertEqual(limits[resource.RLIMIT_CPU], (6, 7))
        self.assertEqual(limits[resource.RLIMIT_FSIZE], (0, 0))
        self.assertEqual(limits[resource.RLIMIT_NOFILE], (32, 32))
        self.assertEqual(limits[resource.RLIMIT_CORE], (0, 0))


class ExtractionProcessTests(unittest.IsolatedAsyncioTestCase):
    async def test_real_isolated_worker_normal_documents_invalid_input_and_env(self):
        create = asyncio.create_subprocess_exec
        environments = []

        async def capture(*args, **kwargs):
            self.assertIn("-I", args)
            self.assertIn("-B", args)
            environments.append(kwargs["env"])
            return await create(*args, **kwargs)

        with patch.object(extraction.asyncio, "create_subprocess_exec", side_effect=capture), patch.dict(os.environ, {"SYNTHETIC_SECRET": "must-not-propagate"}):
            self.assertEqual(await extraction.extract_document("sample.docx", docx_bytes()), "Hello document\nleft | right")
            self.assertEqual(await extraction.extract_document("sample.pdf", pdf_bytes()), "Hello PDF")
            with self.assertRaises(extraction.ExtractionError) as caught:
                await extraction.extract_document("malformed.pdf", b"invalid document: synthetic-secret")
            self.assertEqual(caught.exception.status_code, 422)
            self.assertNotIn("synthetic-secret", caught.exception.detail)
        self.assertTrue(environments)
        for environment in environments:
            self.assertNotIn("SYNTHETIC_SECRET", environment)
            self.assertEqual(set(environment), {"PATH", "LANG", "PYTHONIOENCODING"})

    async def test_timeout_cancellation_saturation_and_recovery(self):
        create = asyncio.create_subprocess_exec
        processes = []
        started = asyncio.Event()

        async def capture(*args, **kwargs):
            process = await create(*args, **kwargs)
            processes.append(process)
            started.set()
            return process

        with tempfile.TemporaryDirectory() as directory:
            script = Path(directory) / "blocked_worker.py"
            script.write_text("import time\ntime.sleep(60)\n")
            with patch.object(extraction, "_WORKER_PATH", script), patch.object(extraction.asyncio, "create_subprocess_exec", side_effect=capture):
                task = asyncio.create_task(extraction.extract_document("slow.pdf", b"input"))
                await asyncio.wait_for(started.wait(), 2)
                # A slow parser leaves the loop responsive and never creates a
                # queued/unbounded worker on saturation.
                await asyncio.sleep(0.01)
                with self.assertRaises(extraction.ExtractionError) as busy:
                    await extraction.extract_document("other.pdf", b"input")
                self.assertEqual(busy.exception.status_code, 503)
                self.assertEqual(len(processes), 1)
                task.cancel()
                with self.assertRaises(asyncio.CancelledError):
                    await task
                self.assertIsNotNone(processes[0].returncode)
                started.clear()
                with patch.object(extraction, "MAX_PROCESSING_SECONDS", 0.05):
                    with self.assertRaises(extraction.ExtractionError) as expired:
                        await extraction.extract_document("slow.pdf", b"input")
                self.assertEqual(expired.exception.status_code, 504)
                self.assertIsNotNone(processes[-1].returncode)
        self.assertEqual(await extraction.extract_document("recovered.pdf", pdf_bytes()), "Hello PDF")

    async def test_repeated_cancellation_during_spawn_retains_handle_and_permit(self):
        create = asyncio.create_subprocess_exec
        spawned, release_spawn = asyncio.Event(), asyncio.Event()
        processes = []

        async def delayed_spawn(*args, **kwargs):
            process = await create(*args, **kwargs)
            processes.append(process)
            spawned.set()
            await release_spawn.wait()
            return process

        with tempfile.TemporaryDirectory() as directory:
            script = Path(directory) / "blocked_worker.py"
            script.write_text("import time\ntime.sleep(60)\n")
            with patch.object(extraction, "_WORKER_PATH", script), patch.object(extraction.asyncio, "create_subprocess_exec", side_effect=delayed_spawn):
                task = asyncio.create_task(extraction.extract_document("slow.pdf", b"input"))
                try:
                    await asyncio.wait_for(spawned.wait(), 2)
                    for _ in range(3):
                        task.cancel()
                        await asyncio.sleep(0)
                    self.assertFalse(task.done())
                    with self.assertRaises(extraction.ExtractionError) as busy:
                        await extraction.extract_document("second.pdf", b"input")
                    self.assertEqual(busy.exception.status_code, 503)
                    self.assertIsNone(processes[0].returncode)
                finally:
                    release_spawn.set()
                    task.cancel()
                    with self.assertRaises(asyncio.CancelledError):
                        await task
                self.assertIsNotNone(processes[0].returncode)
        self.assertEqual(await extraction.extract_document("recovered.pdf", pdf_bytes()), "Hello PDF")

    async def test_repeated_cancellation_during_reaping_retains_permit(self):
        terminate = extraction._terminate
        create = asyncio.create_subprocess_exec
        entered, release_cleanup, spawned = asyncio.Event(), asyncio.Event(), asyncio.Event()
        processes = []

        async def delayed_terminate(process):
            entered.set()
            await release_cleanup.wait()
            await terminate(process)

        async def capture(*args, **kwargs):
            process = await create(*args, **kwargs)
            processes.append(process)
            spawned.set()
            return process

        with tempfile.TemporaryDirectory() as directory:
            script = Path(directory) / "blocked_worker.py"
            script.write_text("import time\ntime.sleep(60)\n")
            for reason in ("caller", "deadline"):
                with self.subTest(reason=reason):
                    entered.clear()
                    release_cleanup.clear()
                    spawned.clear()
                    with patch.object(extraction, "_WORKER_PATH", script), patch.object(extraction, "_terminate", side_effect=delayed_terminate), patch.object(extraction.asyncio, "create_subprocess_exec", side_effect=capture), patch.object(extraction, "MAX_PROCESSING_SECONDS", 0.05):
                        task = asyncio.create_task(extraction.extract_document("slow.pdf", b"input"))
                        try:
                            await asyncio.wait_for(spawned.wait(), 2)
                            if reason == "caller":
                                task.cancel()
                            await asyncio.wait_for(entered.wait(), 2)
                            for _ in range(3):
                                task.cancel()
                                await asyncio.sleep(0)
                            self.assertFalse(task.done())
                            with self.assertRaises(extraction.ExtractionError) as busy:
                                await extraction.extract_document("second.pdf", b"input")
                            self.assertEqual(busy.exception.status_code, 503)
                        finally:
                            release_cleanup.set()
                            task.cancel()
                            with self.assertRaises(asyncio.CancelledError):
                                await task
                        self.assertIsNotNone(processes[-1].returncode)
        self.assertEqual(await extraction.extract_document("recovered.pdf", pdf_bytes()), "Hello PDF")

    async def test_process_creation_failure_is_sanitized_and_releases_slot(self):
        with patch.object(extraction.asyncio, "create_subprocess_exec", side_effect=OSError("synthetic secret")):
            with self.assertRaises(extraction.ExtractionError) as caught:
                await extraction.extract_document("sample.pdf", b"pdf")
        self.assertEqual(caught.exception.status_code, 503)
        self.assertNotIn("secret", caught.exception.detail)
        self.assertEqual(await extraction.extract_document("normal.txt", b"okay"), "okay")
        # Verify a worker slot, rather than only a text path, recovers too.
        self.assertEqual(await extraction.extract_document("normal.pdf", pdf_bytes()), "Hello PDF")

    async def test_bounded_protocol_output_and_non_dictionary_response(self):
        with tempfile.TemporaryDirectory() as directory:
            script = Path(directory) / "protocol_worker.py"
            script.write_text('import sys\nsys.stdout.write("x" * 10_000_000)\n')
            with patch.object(extraction, "_WORKER_PATH", script), patch.object(extraction, "MAX_WIRE_BYTES", 1024):
                with self.assertRaises(extraction.ExtractionError) as caught:
                    await extraction.extract_document("sample.pdf", b"input")
            self.assertEqual(caught.exception.status_code, 413)
            script.write_text('print("[]")\n')
            with patch.object(extraction, "_WORKER_PATH", script):
                with self.assertRaises(extraction.ExtractionError) as caught:
                    await extraction.extract_document("sample.pdf", b"input")
            self.assertEqual(caught.exception.status_code, 503)

    async def test_real_worker_rejects_page_output_and_zip_bombs(self):
        archive = BytesIO()
        with ZipFile(archive, "w", ZIP_DEFLATED) as document:
            document.writestr("large.xml", b"x" * (worker.MAX_EXPANDED_BYTES + 1))
        fixtures = (
            ("many-pages.pdf", pdf_bytes(worker.MAX_PDF_PAGES + 1)),
            ("too-much-text.docx", docx_bytes("a" * (worker.MAX_OUTPUT_CHARACTERS + 1))),
            ("compressed-bomb.docx", archive.getvalue()),
        )
        for name, content in fixtures:
            with self.subTest(name=name):
                self.assertLess(len(content), worker.MAX_INPUT_BYTES)
                with self.assertRaises(extraction.ExtractionError) as caught:
                    await extraction.extract_document(name, content)
                self.assertEqual(caught.exception.status_code, 413)

    async def test_os_cpu_and_memory_enforcement_in_disposable_children(self):
        # Tight synthetic budgets verify OS enforcement without allocating the
        # normal 384 MiB budget or burning the full six-second production CPU
        # allowance. These children contain no document parser or real data.
        scripts = (
            "import resource\nresource.setrlimit(resource.RLIMIT_CPU, (1, 2))\nwhile True: pass\n",
            'import resource,json\nresource.setrlimit(resource.RLIMIT_AS, (64*1024*1024,)*2)\ntry:\n payload=bytearray(128*1024*1024)\nexcept MemoryError:\n print(json.dumps({"ok":False,"kind":"limit"}))\n',
        )
        with tempfile.TemporaryDirectory() as directory:
            script = Path(directory) / "resource_worker.py"
            for source in scripts:
                script.write_text(source)
                with patch.object(extraction, "_WORKER_PATH", script), patch.object(extraction, "MAX_PROCESSING_SECONDS", 4):
                    with self.assertRaises(extraction.ExtractionError) as caught:
                        await extraction.extract_document("sample.pdf", b"input")
                self.assertEqual(caught.exception.status_code, 413)

    async def test_worker_python_network_guard(self):
        with tempfile.TemporaryDirectory() as directory:
            script = Path(directory) / "network_worker.py"
            script.write_text(
                "import runpy,socket,json\n"
                f"worker=runpy.run_path({str(extraction._WORKER_PATH)!r})\n"
                "worker['_disable_network']()\n"
                "try:\n socket.create_connection(('127.0.0.1',1))\n"
                "except OSError:\n print(json.dumps({'ok':True,'text':'blocked'}))\n"
            )
            with patch.object(extraction, "_WORKER_PATH", script):
                self.assertEqual(await extraction.extract_document("sample.pdf", b"input"), "blocked")

    async def test_text_boundaries_and_unicode_replacement_are_preserved(self):
        with patch.object(extraction, "MAX_OUTPUT_CHARACTERS", 3):
            self.assertEqual(await extraction.extract_document("okay.md", b"a\xffb"), "a\ufffdb")
            with self.assertRaises(extraction.ExtractionError) as caught:
                await extraction.extract_document("large.txt", b"abcd")
            self.assertEqual(caught.exception.status_code, 413)


class UploadResourceTests(unittest.IsolatedAsyncioTestCase):
    user = AuthenticatedUser(id="synthetic-owner")

    async def test_all_files_closed_on_first_failure_and_count_limit(self):
        for uploads in (
            [UploadFile(filename="bad.exe", file=BytesIO(b"bad")), UploadFile(filename="later.txt", file=BytesIO(b"valid"))],
            [UploadFile(filename="valid.txt", file=BytesIO(b"x")) for _ in range(6)],
        ):
            with self.assertRaises(HTTPException):
                await files.upload_files(uploads, self.user)
            self.assertTrue(all(upload.file.closed for upload in uploads))

    async def test_aggregate_byte_and_output_boundaries(self):
        for limit_name, limit, expected in (
            ("MAX_AGGREGATE_BYTES", 3, 413),
            ("MAX_AGGREGATE_CHARACTERS", 3, 413),
            ("MAX_AGGREGATE_BYTES", 4, 200),
            ("MAX_AGGREGATE_CHARACTERS", 4, 200),
        ):
            uploads = [UploadFile(filename=f"{index}.txt", file=BytesIO(b"hi")) for index in range(2)]
            with patch.object(files, limit_name, limit):
                if expected == 200:
                    self.assertEqual((await files.upload_files(uploads, self.user))["count"], 2)
                else:
                    with self.assertRaises(HTTPException) as caught:
                        await files.upload_files(uploads, self.user)
                    self.assertEqual(caught.exception.status_code, expected)
            self.assertTrue(all(upload.file.closed for upload in uploads))

    async def test_cancelled_upload_closes_all_files(self):
        entered = asyncio.Event()
        async def blocked(*args):
            entered.set()
            await asyncio.Future()
        uploads = [UploadFile(filename=f"{index}.pdf", file=BytesIO(b"pdf")) for index in range(2)]
        with patch.object(files, "extract_file_text", side_effect=blocked):
            task = asyncio.create_task(files.upload_files(uploads, self.user))
            await asyncio.wait_for(entered.wait(), 2)
            task.cancel()
            with self.assertRaises(asyncio.CancelledError):
                await task
        self.assertTrue(all(upload.file.closed for upload in uploads))
